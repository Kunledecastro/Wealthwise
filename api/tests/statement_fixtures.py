"""Synthetic statement files (no real customer data) in common Nigerian layouts.

Layouts are modelled on the general shape of bank exports (preamble lines, separate
debit/credit columns, newest-first ordering), not copied from any bank.
"""

import io
from datetime import date

# (date, narration, debit, credit, balance_after) oldest → newest. Opening ₦100,000.00.
ROWS = [
    (date(2026, 10, 1), "SALARY OCT 2026 ACME LTD", None, "450,000.00", "550,000.00"),
    (
        date(2026, 10, 2),
        "POS PURCHASE SHOPRITE LEKKI 5399831234567895",
        "25,000.50",
        None,
        "524,999.50",
    ),
    (date(2026, 10, 2), "SMS ALERT CHARGES", "52.50", None, "524,947.00"),
    (date(2026, 10, 3), "NIP TRF TO KUDA/JOHN DOE/0123456789", "100,000.00", None, "424,947.00"),
    (date(2026, 10, 3), "POS PURCHASE BOLT", "500.00", None, "424,447.00"),
    (date(2026, 10, 3), "POS PURCHASE BOLT", "500.00", None, "423,947.00"),  # genuine twin
    (date(2026, 10, 4), "WEB PURCHASE NETFLIX", "4,400.00", None, "419,547.00"),
    (date(2026, 10, 5), "REVERSAL WEB PURCHASE NETFLIX", None, "4,400.00", "423,947.00"),
    (date(2026, 10, 6), "VAT ON CHARGES", "3.94", None, "423,943.06"),
]
OPENING = 10_000_000
CLOSING = 42_394_306


def csv_debit_credit_newest_first() -> bytes:
    lines = [
        "ACCOUNT STATEMENT",
        "Account Name:,JOHN DOE",
        "Period:,01-Oct-2026 to 06-Oct-2026",
        "",
        "Trans. Date,Value Date,Reference,Narration,Debit,Credit,Balance",
    ]
    for i, (d, narr, dr, cr, bal) in enumerate(reversed(ROWS)):
        lines.append(
            f'{d:%d-%b-%Y},{d:%d-%b-%Y},REF{100 + i},"{narr}","{dr or ""}","{cr or ""}","{bal}"'
        )
    lines += ["", "Total,,,,,,"]
    return ("\n".join(lines) + "\n").encode()


def csv_signed_amount() -> bytes:
    lines = ["Date,Description,Amount,Balance"]
    for d, narr, dr, cr, bal in ROWS:
        amt = f"-{dr}" if dr else cr
        lines.append(f'{d:%d/%m/%Y},"{narr}","{amt}","{bal}"')
    return ("\n".join(lines) + "\n").encode()


def csv_indicator_semicolon() -> bytes:
    lines = ["Transaction Date;Details;Amount (NGN);DR/CR"]
    for d, narr, dr, cr, _ in ROWS:
        lines.append(f"{d:%Y-%m-%d};{narr};{dr or cr};{'DR' if dr else 'CR'}")
    return ("\n".join(lines) + "\n").encode()


def _money(s: str | None) -> float | None:
    return None if s is None else float(s.replace(",", ""))


def xlsx_money_in_out() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["Wallet statement"])
    ws.append([])
    ws.append(["Date", "Description", "Money In", "Money Out", "Balance"])
    for d, narr, dr, cr, bal in ROWS:
        ws.append([d, narr, _money(cr), _money(dr), _money(bal)])
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Date", "Amount"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xls_debit_credit() -> bytes:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    datefmt = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, h in enumerate(["Tran Date", "Remarks", "Withdrawals", "Lodgements", "Balance"]):
        ws.write(0, c, h)
    for r, (d, narr, dr, cr, bal) in enumerate(ROWS, start=1):
        ws.write(r, 0, d, datefmt)
        ws.write(r, 1, narr)
        if dr:
            ws.write(r, 2, _money(dr))
        if cr:
            ws.write(r, 3, _money(cr))
        ws.write(r, 4, _money(bal))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
